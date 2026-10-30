#!/usr/bin/env python3
"""Isi nomor halaman daftar isi buku hasil assembly (Mesin 2).

Latar (27 Okt 2026): `buku_umum_assembly.py` menulis daftar isi dengan
tab ber-dot-leader TAPI tanpa angka — hanya "Kata Pengantar" yang punya
"iii". Akibatnya daftar isi buku produksi kosong nomornya (terbukti di
kode dan di buku nyata). Mesin 1 (template-buku-resmi) TIDAK disentuh
untuk ini: ia tidak dipanggil app dan tidak ada di image Docker, jadi
perbaikan di sana = nol efek ke pelanggan.

Cara kerja (jalur operator, mesin yang sudah ada saja):
  1. Render DOCX -> PDF lewat WordCraft (`perender_halaman`; BUKAN
     soffice — soffice dihapus dari desain 26 Okt 2026).
  2. Baca teks per halaman lewat `pdftotext` (di hermes: salinan
     permanen /opt/data/services/pdftotext/bin — JANGAN /tmp).
  3. Cari halaman tiap judul, LEWATI SELURUH span halaman daftar isi
     (pelajaran perbaikan span TOC: saringan titik saja rapuh lintas
     perender karena judul membungkus ke baris tanpa titik).
  4. Tulis angka sebagai run baru sesudah run tab di paragraf daftar
     isi. Entri yang sudah bernomor dilewati (idempotent: jalan dua
     kali = hasil sama).

  Aturan heading-match (G5, 30 Okt 2026): entri TOC HANYA dipetakan
  ke paragraf heading betulan di badan, bukan ke paragraf pertama
  yang teksnya cocok. Syaratnya tiga lapis: (a) di DOCX harus ada
  paragraf badan sesudah blok TOC yang teks normalisasinya PERSIS
  sama dengan heading entri (untuk 'BAB N. Judul' headingnya adalah
  paragraf 'BAB N'; untuk 'Indeks' aliasnya 'INDEKS SUBJEK' lalu
  'INDEKS NAMA'); tanpa heading seperti itu entri masuk HILANG,
  tidak pernah diisi dari teks badan; (b) di render, halaman dicari
  secara terurut (posisi pindai hanya maju, jadi nomor TOC monoton
  dan paragraf badan sebelum heading sebelumnya tidak bisa membajak
  entri sesudahnya); (c) baris render harus SAMA PERSIS dengan
  heading (sesudah normalisasi) atau merupakan potongan awal heading
  yang dilengkapi baris-baris berikutnya sampai persis sama (judul
  panjang yang membungkus). Awalan saja TIDAK cukup: baris badan
  'Daftar Pustaka induk wajib ...' lebih panjang dari heading
  'DAFTAR PUSTAKA', jadi tidak pernah cocok sebagai heading itu.
  Saat mengisi nomor, paragraf TOC dijamin punya karakter tab dan
  tab stop rata kanan 15,5 cm ber-dot-leader, jadi nomor tidak pernah
  menempel pada judul walau entri asalnya tanpa tab.
  5. Simpan ke KELUAR. Angka = nomor halaman hasil render (bukan nomor
     folio cetak — bila buku mencetak folio berbeda, itu soal terpisah).

Pakai: python3 isi_nomor_toc.py MASUK [KELUAR]
  KELUAR default = MASUK dengan ".bernomor" sebelum ".docx".
Keluar: 0 bila semua terisi; 1 bila ada yang hilang (berkas tetap
ditulis dengan yang terisi + daftar HILANG dicetak lantang); 2 bila
gagal total (tidak bisa render/baca/halaman tidak ditemukan di berkas).
"""

import re
import subprocess
import sys
import tempfile
from pathlib import Path

import perender_halaman
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_LEADER
from docx.shared import Cm


def halaman_teks(pdf, total):
    """Teks per halaman (huruf besar), via pdftotext -layout."""
    out = []
    for i in range(1, total + 1):
        t = subprocess.run(['pdftotext', '-f', str(i), '-l', str(i),
                            '-layout', str(pdf), '-'],
                           check=True, capture_output=True,
                           text=True).stdout
        out.append(t.upper())
    return out


def span_toc(halaman):
    """Halaman daftar isi + lanjutan berurutan bertitik 10+ (1-based)."""
    mulai = None
    for i, t in enumerate(halaman):
        if any(b.strip() == 'DAFTAR ISI' for b in t.splitlines()):
            mulai = i + 1
            break
    if mulai is None:
        return set()
    span = {mulai}
    h = mulai + 1
    while h <= len(halaman):
        n = sum(1 for b in halaman[h - 1].splitlines()
                if re.search(r'\.{10,}', b.strip()))
        if n < 3:
            break
        span.add(h)
        h += 1
    return span


def normalisasi_baris(teks):
    """Normalisasi satu baris untuk pencocokan: huruf besar, spasi ciut."""
    return ' '.join(teks.split()).upper()


def heading_dari_entri(judul):
    """Daftar heading badan yang sah untuk satu entri TOC (berurutan).

    'BAB N. Judul' -> paragraf pembuka 'BAB N' (judul babnya ada di
    paragraf berikutnya, tapi halaman bab = halaman pembuka itu).
    Sub-bab 'N.M Judul' -> teks entri itu sendiri. Back matter ->
    teks entri itu sendiri, dengan alias: 'Indeks' boleh dipetakan ke
    'INDEKS SUBJEK' (didahulukan) atau 'INDEKS NAMA', karena assembly
    tidak pernah menulis heading tunggal 'INDEKS'.
    """
    atas = normalisasi_baris(judul)
    m = re.match(r'^BAB\s+(\d+)\.', atas)
    if m:
        return ['BAB ' + m.group(1)]
    if atas == 'INDEKS':
        return ['INDEKS SUBJEK', 'INDEKS NAMA']
    return [atas]


def kunci_dari_judul(judul):
    """Kunci cari judul di halaman isi (kompatibilitas: heading pertama)."""
    return heading_dari_entri(judul)[0]


def cari_heading_docx(doc, akhir, judul):
    """Heading badan untuk satu entri TOC, atau None bila tidak ada.

    Hanya paragraf sesudah blok TOC yang teks normalisasinya persis
    sama dengan salah satu heading sah entri itu. Paragraf badan
    yang HANYA menyebut judul (misal kalimat '... Daftar Pustaka
    induk wajib ...') tidak pernah lolos, karena teksnya adalah
    kalimat panjang, bukan persis judul itu.
    """
    if akhir is None:
        return None
    for calon in heading_dari_entri(judul):
        for p in doc.paragraphs[akhir:]:
            if normalisasi_baris(p.text) == calon:
                return calon
    return None


def _cocok_heading_di_halaman(baris_halaman, mulai, heading):
    """Jumlah baris yang dipakai heading mulai baris `mulai`, atau 0.

    Baris-baris kosong dilewati. Cocok bila gabungan baris-baris
    (dinormalisasi, dipisah satu spasi) PERSIS sama dengan heading.
    Satu baris yang LEBIH PANJANG dari heading (misal baris badan
    'DAFTAR PUSTAKA INDUK ...') tidak pernah cocok, karena gabungan
    itu bukan persis heading. Judul panjang yang membungkus ke
    baris berikut tetap cocok lewat gabungan beberapa baris.
    """
    gabung = []
    for k in range(mulai, len(baris_halaman)):
        b = normalisasi_baris(baris_halaman[k])
        if not b:
            continue
        gabung.append(b)
        teks = ' '.join(gabung)
        if teks == heading:
            return k - mulai + 1
        if not heading.startswith(teks):
            return 0
        if len(gabung) >= 4:
            return 0
    return 0


def cari_halaman_terurut(heading_list, halaman, toc_pages):
    """Halaman tiap heading, terurut dan monoton (daftar sejajar input).

    Pindai hanya maju: pencarian heading ke-n dimulai dari posisi
    (halaman, baris) tempat heading ke-(n-1) ditemukan. Akibatnya
    nomor TOC tidak pernah turun, dan paragraf badan SEBELUM heading
    sebelumnya tidak bisa membajak entri sesudahnya. Heading yang
    tidak ditemukan -> None (pemanggil memasukkannya ke HILANG,
    bukan mengarang angka dari teks badan). Bila satu heading gagal
    ditemukan, pindai heading berikutnya dimulai lagi dari posisi
    awal pencarian heading yang gagal itu, supaya satu kegagalan
    tidak menghanguskan semua entri sesudahnya.
    """
    baris_per_halaman = [t.splitlines() for t in halaman]
    hasil = []
    hal = 0
    bar = 0
    for heading in heading_list:
        hal_awal, bar_awal = hal, bar
        ketemu = None
        while hal < len(baris_per_halaman):
            h = hal + 1
            if h in toc_pages:
                hal += 1
                bar = 0
                continue
            baris = baris_per_halaman[hal]
            pakai = 0
            while bar < len(baris):
                pakai = _cocok_heading_di_halaman(baris, bar, heading)
                if pakai:
                    break
                bar += 1
            if pakai:
                ketemu = h
                bar = bar + pakai
                if bar >= len(baris):
                    hal += 1
                    bar = 0
                break
            hal += 1
            bar = 0
        if ketemu is None:
            hal, bar = hal_awal, bar_awal
        hasil.append(ketemu)
    return hasil


def cari_halaman(kunci, halaman, toc_pages):
    """Halaman pertama (di luar span TOC) yang barisnya PERSIS = kunci.

    Dipakai kompatibilitas saja; main() memakai cari_halaman_terurut.
    Awalan sengaja TIDAK dipakai lagi: baris 'DAFTAR PUSTAKA INDUK
    ...' tidak boleh cocok sebagai 'DAFTAR PUSTAKA'.
    """
    kunci = normalisasi_baris(kunci)
    for i, t in enumerate(halaman):
        h = i + 1
        if h in toc_pages:
            continue
        for baris in t.splitlines():
            if normalisasi_baris(baris) == kunci:
                return h
    return None


def _sudah_bernomor(paragraf):
    """True bila paragraf TOC sudah punya run angka sesudah tab."""
    runs = paragraf.runs
    if not runs:
        return False
    ada_tab = any('\t' in r.text for r in runs)
    terakhir = runs[-1].text.strip()
    return ada_tab and re.fullmatch(r'[ivxlcdm\d]+', terakhir, re.I) is not None


def paragraf_toc(doc):
    """(entri, sudah, akhir); entri = [(judul, paragraf)] yang perlu diisi.

    akhir = indeks paragraf tempat blok TOC berakhir (paragraf
    page-break penutup TOC, atau paragraf 'BAB 1' pertama di badan,
    atau jumlah paragraf bila tidak ada pembatas). Pencarian heading
    badan HARUS mulai dari akhir (inklusif), supaya entri TOC sendiri
    tidak pernah dicocokkan sebagai headingnya (teks entri sub-bab
    dan back matter sama persis dengan teks heading badannya).
    """
    paras = doc.paragraphs
    awal = None
    for i, p in enumerate(paras):
        if p.text.strip().upper() == 'DAFTAR ISI':
            awal = i
            break
    if awal is None:
        return None, None, None
    entri, sudah = [], []
    akhir = len(paras)
    for idx in range(awal + 1, len(paras)):
        p = paras[idx]
        xml = p._p.xml
        if 'w:type="page"' in xml:
            akhir = idx
            break
        if p.text.strip() == 'BAB 1':
            akhir = idx
            break
        judul = p.text.strip()
        if not judul:
            continue
        if _sudah_bernomor(p):
            sudah.append((judul.split('\t')[0].strip(),
                          p.runs[-1].text.strip()))
        else:
            entri.append((judul, p))
    return entri, sudah, akhir


def _salin_font(run_baru, run_contoh):
    try:
        run_baru.font.name = run_contoh.font.name
        run_baru.font.size = run_contoh.font.size
        run_baru.font.bold = run_contoh.font.bold
        run_baru.font.italic = run_contoh.font.italic
    except Exception:
        pass


def _pastikan_tab(p):
    """Jamin paragraf TOC punya tab + tab stop rata kanan 15,5 cm.

    Bila entri asalnya tanpa tab (atau tanpa tab stop), nomor yang
    ditambahkan sesudahnya akan menempel pada judul. Fungsi ini
    menambahkan run tab (bila belum ada) dan tab stop rata kanan
    ber-dot-leader (bila belum ada di posisi itu).
    """
    if not any('\t' in r.text for r in p.runs):
        contoh = p.runs[0] if p.runs else None
        tab = p.add_run('\t')
        if contoh is not None:
            _salin_font(tab, contoh)
    for ts in p.paragraph_format.tab_stops:
        try:
            posisi = ts.position.cm
        except Exception:
            continue
        if abs(posisi - 15.5) < 0.05:
            return
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT, WD_TAB_LEADER.DOTS)


def main(argv):
    if len(argv) < 2 or len(argv) > 3:
        print('Pakai: python3 isi_nomor_toc.py MASUK [KELUAR]')
        return 2
    masuk = Path(argv[1])
    if len(argv) == 3:
        keluar = Path(argv[2])
    else:
        keluar = masuk.with_name(masuk.stem + '.bernomor' + masuk.suffix)
    if not masuk.exists():
        print(f'GAGAL: berkas tidak ada: {masuk}')
        return 2

    total = perender_halaman.hitung_halaman(str(masuk))
    if total is None:
        print('GAGAL: perender tidak tersedia (WORDCRAFT_CLI/PATH/docker hermes).')
        return 2
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / 'buku.pdf'
        if not perender_halaman.render_pdf(str(masuk), str(pdf)):
            print('GAGAL: render PDF gagal.')
            return 2
        try:
            halaman = halaman_teks(pdf, total)
        except FileNotFoundError:
            print('GAGAL: pdftotext tidak ditemukan — di hermes pakai '
                  'PATH=/opt/data/services/pdftotext/bin:$PATH.')
            return 2
        toc_pages = span_toc(halaman)

        doc = Document(str(masuk))
        entri, sudah, akhir = paragraf_toc(doc)
        if entri is None:
            print('GAGAL: paragraf DAFTAR ISI tidak ditemukan di berkas.')
            return 2
        rencana, hilang = [], []
        for judul, p in entri:
            heading = cari_heading_docx(doc, akhir, judul)
            if heading is None:
                hilang.append(judul)
                continue
            rencana.append((judul, p, heading))
        peta = cari_halaman_terurut([r[2] for r in rencana],
                                     halaman, toc_pages)
        terisi = []
        for (judul, p, _heading), h in zip(rencana, peta):
            if h is None:
                hilang.append(judul)
                continue
            _pastikan_tab(p)
            contoh = p.runs[0] if p.runs else None
            baru = p.add_run(str(h))
            if contoh is not None:
                _salin_font(baru, contoh)
            terisi.append((judul, h))
        doc.save(str(keluar))

    for judul, h in terisi:
        print(f'  TERISI: {judul} -> {h}')
    for judul, _ in sudah:
        print(f'  SUDAH: {judul} (dibiarkan)')
    if hilang:
        print('HILANG (tidak ditulis, periksa judulnya):')
        for judul in hilang:
            print(f'  - {judul}')
    print(f'RINGKAS ISI NOMOR TOC: {len(terisi)} terisi, '
          f'{len(sudah)} sudah, {len(hilang)} hilang -> {keluar}')
    return 1 if hilang else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))

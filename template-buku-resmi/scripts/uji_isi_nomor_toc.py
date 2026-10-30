#!/usr/bin/env python3
"""Uji regresi isi_nomor_toc.py pada buku sintetis gaya assembly.

Buku sintetis (page break antar bagian = halaman deterministik di
perender mana pun): 1 sampul, 2 kata pengantar, 3 daftar isi, 4 BAB 1,
5 sub 1.1, 6 BAB 2, 7 glosarium. Daftar isi meniru add_daftar_isi:
judul + tab ber-dot-leader tanpa angka (kecuali Kata Pengantar "iii").

Kasus: (a) angka terisi persis 4/5/6/7 dan BUKAN halaman TOC (3) —
regresi utama: judul di halaman TOC tidak boleh terpetakan ke TOC;
(b) "iii" dibiarkan; (c) jalan kedua kali = angka sama (idempotent).
SKIP (exit 0) bila perender/pdftotext tidak tersedia.
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_LEADER
from docx.shared import Cm, Pt

import perender_halaman

HERE = Path(__file__).resolve().parent


def _entri(doc, teks, tab=True, ukuran=None, tebal=False, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    r = p.add_run(teks)
    r.font.name = 'Times New Roman'
    if ukuran is not None:
        r.font.size = Pt(ukuran)
    r.bold = tebal
    if tab:
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run('\t')
    return p


def bangun_buku(path):
    doc = Document()
    doc.add_paragraph('SAMPUL')
    doc.add_page_break()
    doc.add_paragraph('KATA PENGANTAR')
    doc.add_paragraph('Isi kata pengantar.')
    doc.add_page_break()
    h = doc.add_paragraph('DAFTAR ISI')
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _entri(doc, 'Kata Pengantar')
    doc.paragraphs[-1].add_run('iii')
    _entri(doc, 'BAB 1. Mengamati Langit', ukuran=12, tebal=True)
    _entri(doc, '1.1 Matahari, Suhu, dan Bayangan', ukuran=11, indent=0.8)
    _entri(doc, 'BAB 2. Saat Cuaca Berubah', ukuran=12, tebal=True)
    _entri(doc, 'Glosarium')
    doc.add_page_break()
    doc.add_paragraph('BAB 1')
    doc.add_paragraph('Isi bab satu.')
    doc.add_page_break()
    doc.add_paragraph('1.1 Matahari, Suhu, dan Bayangan')
    doc.add_paragraph('Isi sub bab.')
    doc.add_page_break()
    doc.add_paragraph('BAB 2')
    doc.add_paragraph('Isi bab dua.')
    doc.add_page_break()
    doc.add_paragraph('GLOSARIUM')
    doc.add_paragraph('Istilah penting.')
    doc.save(str(path))


def bangun_buku_jebakan(path):
    """Buku jebakan first-match: paragraf badan yang menyebut judul
    back matter SEBELUM heading back matter yang asli.

    Halaman: 1 sampul, 2 kata pengantar, 3 daftar isi, 4 BAB 1,
    5 paragraf badan 'Daftar Pustaka induk ...' (satu baris, diawali
    persis 'Daftar Pustaka '), 6 DAFTAR PUSTAKA, 7 TENTANG PENULIS.
    Kode lama (startswith) memetakan Daftar Pustaka ke 5; kode baru
    (heading-match) harus memetakannya ke 6. Entri TOC Daftar
    Pustaka sengaja dibuat TANPA tab untuk menguji jaminan tab.
    """
    doc = Document()
    doc.add_paragraph('SAMPUL')
    doc.add_page_break()
    doc.add_paragraph('KATA PENGANTAR')
    doc.add_paragraph('Isi kata pengantar.')
    doc.add_page_break()
    h = doc.add_paragraph('DAFTAR ISI')
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _entri(doc, 'Kata Pengantar')
    doc.paragraphs[-1].add_run('iii')
    _entri(doc, 'BAB 1. Mengamati Langit', ukuran=12, tebal=True)
    _entri(doc, 'Daftar Pustaka', tab=False)
    _entri(doc, 'Tentang Penulis')
    doc.add_page_break()
    doc.add_paragraph('BAB 1')
    doc.add_paragraph('Isi bab satu.')
    doc.add_page_break()
    doc.add_paragraph('Daftar Pustaka induk wajib diverifikasi '
                      'sebelum naskah dikunci.')
    doc.add_page_break()
    doc.add_paragraph('DAFTAR PUSTAKA')
    doc.add_paragraph('Alexandrova E. (2021). Contoh referensi.')
    doc.add_page_break()
    doc.add_paragraph('TENTANG PENULIS')
    doc.add_paragraph('Bio penulis.')
    doc.save(str(path))


def toc_terakhir(doc):
    """judul -> run terakhir, untuk paragraf sesudah DAFTAR ISI."""
    paras = doc.paragraphs
    awal = next(i for i, p in enumerate(paras)
                if p.text.strip().upper() == 'DAFTAR ISI')
    hasil = {}
    for p in paras[awal + 1:]:
        if 'w:type="page"' in p._p.xml or p.text.strip() == 'BAB 1':
            break
        if not p.runs:
            continue
        judul = p.runs[0].text.strip()
        if judul:
            hasil[judul] = p.runs[-1].text.strip()
    return hasil


def main():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        masuk, keluar, keluar2 = (tmp / 'buku.docx', tmp / 'buku1.docx',
                                  tmp / 'buku2.docx')
        bangun_buku(masuk)
        if shutil.which('pdftotext') is None:
            print('SKIP: pdftotext tidak tersedia')
            return 0
        if perender_halaman.hitung_halaman(str(masuk)) is None:
            print('SKIP: perender tidak tersedia')
            return 0
        r = subprocess.run([sys.executable, str(HERE / 'isi_nomor_toc.py'),
                            str(masuk), str(keluar)],
                           capture_output=True, text=True)
        print(r.stdout, end='')
        if r.returncode != 0:
            print(f'GAGAL: filler exit {r.returncode}\n{r.stderr}')
            return 1
        toc = toc_terakhir(Document(str(keluar)))
        harap = {'BAB 1. Mengamati Langit': '4',
                 '1.1 Matahari, Suhu, dan Bayangan': '5',
                 'BAB 2. Saat Cuaca Berubah': '6',
                 'Glosarium': '7'}
        gagal = []
        for judul, angka in harap.items():
            got = toc.get(judul)
            tag = 'LULUS' if got == angka else 'GAGAL'
            print(f'{tag} [{judul}]: {got} (harap {angka})')
            if got != angka:
                gagal.append(judul)
        iii = toc.get('Kata Pengantar')
        print(f'{"LULUS" if iii == "iii" else "GAGAL"} [Kata Pengantar '
              f'dibiarkan]: {iii}')
        if iii != 'iii':
            gagal.append('Kata Pengantar')
        r2 = subprocess.run([sys.executable, str(HERE / 'isi_nomor_toc.py'),
                             str(keluar), str(keluar2)],
                            capture_output=True, text=True)
        toc2 = toc_terakhir(Document(str(keluar2)))
        sama = all(toc2.get(j) == a for j, a in harap.items())
        print(f'{"LULUS" if sama and r2.returncode == 0 else "GAGAL"} '
              f'[idempotent, jalan kedua]: {toc2}')
        if not (sama and r2.returncode == 0):
            gagal.append('idempotent')
        # Kasus jebakan first-match + jaminan tab (G5).
        jebakan_masuk, jebakan_keluar = (tmp / 'jebakan.docx',
                                         tmp / 'jebakan1.docx')
        bangun_buku_jebakan(jebakan_masuk)
        rj = subprocess.run([sys.executable, str(HERE / 'isi_nomor_toc.py'),
                             str(jebakan_masuk), str(jebakan_keluar)],
                            capture_output=True, text=True)
        print(rj.stdout, end='')
        if rj.returncode != 0:
            print(f'GAGAL: filler jebakan exit {rj.returncode}\n{rj.stderr}')
            gagal.append('jebakan exit')
        else:
            tocj = toc_terakhir(Document(str(jebakan_keluar)))
            for judul, angka in [('Daftar Pustaka', '6'),
                                 ('Tentang Penulis', '7')]:
                got = tocj.get(judul)
                tag = 'LULUS' if got == angka else 'GAGAL'
                print(f'{tag} [jebakan {judul}]: {got} (harap {angka})')
                if got != angka:
                    gagal.append('jebakan ' + judul)
            dj = Document(str(jebakan_keluar))
            parasj = dj.paragraphs
            awalj = next(i for i, p in enumerate(parasj)
                         if p.text.strip().upper() == 'DAFTAR ISI')
            pj = next(p for p in parasj[awalj + 1:]
                      if p.runs and p.runs[0].text.strip() == 'Daftar Pustaka')
            ada_tab = '\t' in pj.text
            ada_stop = any(ts.alignment == WD_ALIGN_PARAGRAPH.RIGHT
                           for ts in pj.paragraph_format.tab_stops)
            print(f'{"LULUS" if ada_tab else "GAGAL"} '
                  f'[jebakan tab ditambahkan]: {pj.text!r}')
            print(f'{"LULUS" if ada_stop else "GAGAL"} '
                  f'[jebakan tab stop rata kanan]: '
                  f'{[(round(ts.position.cm, 1), int(ts.alignment)) for ts in pj.paragraph_format.tab_stops]}')
            if not ada_tab:
                gagal.append('jebakan tab')
            if not ada_stop:
                gagal.append('jebakan tab stop')
        print(f'RINGKAS UJI ISI NOMOR TOC: '
              f'{len(harap) + 2 + 4 - len(gagal)} LULUS / {len(gagal)} GAGAL')
        return 1 if gagal else 0


if __name__ == '__main__':
    sys.exit(main())

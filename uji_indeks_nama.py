"""Uji perbaikan indeks nama + cari judul (akar R2 job 220), tanpa LLM/API.
Jalankan dari checkout: /opt/data/.venv/bin/python3 uji_indeks_nama.py
Keluar 0 bila semua LULUS, 1 bila ada GAGAL.
"""

import os
import sys
import tempfile

for p in (os.getcwd(), '/opt/data/services/eduai-publishing', '/app'):
    if p and os.path.isdir(p) and p not in sys.path:
        sys.path.insert(0, p)

from app import referensi_indeks as RI

LULUS = 0
GAGAL = []


def cek(nama, ok, info=''):
    global LULUS
    if ok:
        LULUS += 1
        print('  LULUS %s' % nama)
    else:
        GAGAL.append(nama)
        print('  GAGAL %s %s' % (nama, info))


def test_kumpulkan_nama_apa():
    hasil = RI.kumpulkan_nama([
        'Yang, H. (2020). Judul.',
        'Smith, J. P. (2021). Judul.',
    ])
    cek('nama-apa', hasil == ['Yang', 'Smith'], repr(hasil))


def test_kumpulkan_nama_depan_belakang():
    hasil = RI.kumpulkan_nama([
        'Wan Laura Hardilawati (2020). Judul.',
        'Chege Samwel Macharia (2020). Judul.',
        'Wang Daoping (2020). Judul.',
        'Dabić Marina (2019). Judul.',
        'Prasanna RPIR (2021). Judul.',
        'Hervás-Oliver José-Luis (2018). Judul.',
    ])
    cek('nama-depan-belakang-surname',
        'Hardilawati' in hasil and 'Macharia' in hasil
        and 'Daoping' in hasil and 'Prasanna' in hasil,
        repr(hasil))
    cek('nama-depan-belakang-bukan-frasa-lengkap',
        'Wan Laura Hardilawati' not in hasil
        and 'Chege Samwel Macharia' not in hasil,
        repr(hasil))
    cek('nama-depan-belakang-bukan-inisial',
        'RPIR' not in hasil, repr(hasil))


def test_kumpulkan_nama_institusi_inisial():
    hasil = RI.kumpulkan_nama([
        'Universitas Tidaar, et al. (2020). Judul.',
        'Lu Zhiqiang (2020). Judul.',
        'Karami Masoud (2020). Judul.',
    ])
    cek('nama-institusi-inisial',
        hasil == ['Lu', 'Zhiqiang', 'Karami', 'Masoud'], repr(hasil))


def test_halaman_cocok():
    halaman = {1: 'BAB 1\nKutipan Hardilawati (2020).',
               2: 'GLOSARIUM\nDaftar istilah.'}
    isi = {1: halaman[1]}
    cek('cocok-frasa-lengkap-kosong',
        RI._halaman_cocok(halaman, isi, 'Wan Laura Hardilawati') == [],
        repr(RI._halaman_cocok(halaman, isi, 'Wan Laura Hardilawati')))
    cek('cocok-surname-satu-halaman',
        RI._halaman_cocok(halaman, isi, 'Hardilawati') == [1],
        repr(RI._halaman_cocok(halaman, isi, 'Hardilawati')))


def test_isi_indeks_trap():
    from docx import Document
    doc = Document()
    for teks in [
        'BAB 1 Transformasi Digital UMKM',
        'Badan bab satu mengutip Hardilawati (2020).',
        'Indeks',
        'KALIMAT_PENANDA_BODY_SETELAH_INDEKS',
        'GLOSARIUM',
        'INDEKS SUBJEK',
        '(belum ada entri, isi setelah naskah selesai)',
        'INDEKS NAMA',
        'DAFTAR PUSTAKA',
    ]:
        doc.add_paragraph(teks)
    tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
    tmp.close()
    doc.save(tmp.name)
    teks_halaman = {
        1: 'BAB 1\nKutipan Hardilawati (2020) soal UMKM dan Adopsi teknologi.',
        2: 'GLOSARIUM\nDaftar istilah.',
        3: 'INDEKS SUBJEK\nIndeks subjek.',
        4: 'DAFTAR PUSTAKA\nDaftar pustaka.',
    }
    nama_list = RI.kumpulkan_nama([
        'Wan Laura Hardilawati (2020). Judul.',
        'Chege Samwel Macharia (2020). Judul.',
    ])
    hasil = RI.isi_indeks(
        tmp.name, teks_halaman,
        {'UMKM': 'Usaha Mikro Kecil Menengah',
         'Adopsi teknologi': 'penerapan teknologi'},
        nama_list)
    sesudah = Document(tmp.name)
    semua = [p.text for p in sesudah.paragraphs]

    def seksi(judul):
        try:
            i = semua.index(judul)
        except ValueError:
            return []
        out = []
        for t in semua[i + 1:]:
            if t.strip().upper() in RI._SEKSI:
                break
            out.append(t)
        return out

    subjek = seksi('INDEKS SUBJEK')
    nama = seksi('INDEKS NAMA')
    cek('indeks-trap-subjek-terisi',
        any('UMKM' in t for t in subjek)
        and any('Adopsi teknologi' in t for t in subjek),
        repr(subjek))
    cek('indeks-trap-nama-surname',
        any('Hardilawati' in t for t in nama),
        repr(nama))
    cek('indeks-trap-nama-bukan-nama-depan',
        not any(t == 'Wan' or t.startswith('Wan\t') for t in nama),
        repr(nama))
    cek('indeks-trap-placeholder-hilang',
        '(belum ada entri, isi setelah naskah selesai)' not in semua,
        repr(semua))
    cek('indeks-trap-sentinel-selamat',
        'KALIMAT_PENANDA_BODY_SETELAH_INDEKS' in semua,
        repr(semua))
    cek('indeks-trap-baris-indeks-selamat',
        semua.count('Indeks') == 1,
        repr([t for t in semua if t == 'Indeks']))
    cek('indeks-trap-hasil',
        hasil['subjek'] >= 2 and hasil['nama'] >= 1, repr(hasil))
    os.unlink(tmp.name)


def test_sapuan():
    try:
        from buku_umum_assembly import bersihkan_em_dash_dokumen
    except Exception as e:
        cek('sapuan-tanpa-em-dash', False, 'impor gagal: %s' % e)
        return
    from docx import Document
    doc = Document()
    doc.add_paragraph('Kalimat dengan ' + chr(0x2014) + ' sisipan.')
    tabel = doc.add_table(rows=1, cols=1)
    tabel.cell(0, 0).paragraphs[0].add_run('Sel ' + chr(0x2014) + ' tabel.')
    tmp = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
    tmp.close()
    doc.save(tmp.name)
    doc2 = Document(tmp.name)
    bersihkan_em_dash_dokumen(doc2)
    doc2.save(tmp.name)
    sesudah = Document(tmp.name)
    sisa_para = sum(p.text.count(chr(0x2014)) for p in sesudah.paragraphs)
    sisa_tabel = 0
    for tb in sesudah.tables:
        for baris in tb.rows:
            for sel in baris.cells:
                for p in sel.paragraphs:
                    sisa_tabel += p.text.count(chr(0x2014))
    cek('sapuan-tanpa-em-dash', sisa_para == 0 and sisa_tabel == 0,
        'para=%d tabel=%d' % (sisa_para, sisa_tabel))
    os.unlink(tmp.name)


if __name__ == '__main__':
    test_kumpulkan_nama_apa()
    test_kumpulkan_nama_depan_belakang()
    test_kumpulkan_nama_institusi_inisial()
    test_halaman_cocok()
    test_isi_indeks_trap()
    test_sapuan()
    total = LULUS + len(GAGAL)
    print('%d/%d LULUS' % (LULUS, total))
    if GAGAL:
        print('GAGAL: %s' % ', '.join(GAGAL))
        sys.exit(1)
    sys.exit(0)

"""Uji sapuan dokumen anti em dash (U+2014), 29 Okt 2026. Jalankan dari root:
  /opt/data/.venv/bin/python3 uji_em_dash_dokumen.py
Berkas ini tidak mengandung karakter U+2014 literal; semua memakai escape."""

import sys

from docx import Document

from buku_umum_assembly import bersihkan_em_dash_dokumen

EM = '\u2014'

doc = Document()
doc.add_paragraph('Bab ini membahas teori \u2014 dari dasar sampai penerapan.')
doc.add_paragraph('Judul Bab \u2014 Konsep Dasar')
tabel = doc.add_table(rows=1, cols=1)
tabel.cell(0, 0).text = 'Istilah \u2014 definisi singkat.'
kaki = doc.sections[0].footer.paragraphs
if kaki:
    kaki[0].text = 'Footer \u2014 Nama Penulis'
bersihkan_em_dash_dokumen(doc)


def semua_teks(d):
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                out += [p.text for p in c.paragraphs]
    for s in d.sections:
        for b in (s.header, s.footer):
            out += [p.text for p in b.paragraphs]
    return out


teks_list = semua_teks(doc)
ok = all(EM not in t for t in teks_list)
print(('LULUS' if ok else 'GAGAL'), '- tidak ada U+2014 di dokumen |', teks_list)
ok2 = (teks_list[0] == 'Bab ini membahas teori, dari dasar sampai penerapan.')
print(('LULUS' if ok2 else 'GAGAL'), '- koma sebagai pengganti |', repr(teks_list[0]))
with open('buku_umum_assembly.py', encoding='utf-8') as f:
    sumber = f.read()
awal = sumber.find('pengantar = f"""')
akhir = sumber.find('"""', awal + len('pengantar = f"""'))
template = sumber[awal:akhir]
ok3 = (EM not in template)
print(('LULUS' if ok3 else 'GAGAL'),
      '- template Kata Pengantar tanpa U+2014')
semua = ok and ok2 and ok3
print('UJI-EM-DASH-DOKUMEN', 'LULUS SEMUA' if semua else 'ADA GAGAL')
sys.exit(0 if semua else 1)

"""Regresi app/referensi_indeks.py — jalan TANPA pytest (stdlib + python-docx).

Jalankan dari root repo:  python3 tests/test_referensi_indeks.py

BUG 28 Okt 2026 (temuan review agen, terbukti): isi_indeks crash TypeError
saat docx BELUM punya seksi 'Indeks Nama' — bentuk normal produksi, karena
fallback add_indeks di kontainer hanya menulis 'INDEKS' (tanpa Indeks Nama).
Akar: cari_judul membaca snapshot `paras = doc.paragraphs` yang tidak melihat
judul yang baru dibuat buat_judul(). Tes test_tanpa_indeks_nama + 
test_tanpa_seksi_indeks_sama_sekali menjaga jalur pembuatan judul itu;
sebelum perbaikan, keduanya crash TypeError.
"""

import json
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'app'))
import referensi_indeks as RI  # noqa: E402
from docx import Document  # noqa: E402


def _docx(path, paragraf):
    doc = Document()
    for t in paragraf:
        doc.add_paragraph(t)
    doc.save(path)
    return path


def _teks(path):
    return '\n'.join(p.text for p in Document(path).paragraphs)


# Halaman: 1 sampul, 2 TOC (menyebut istilah!), 3-5 isi, 6+ back-matter.
HALAMAN = [
    'Sampul',
    'DAFTAR ISI\nPembelian impulsif ....',
    'BAB 1\nBab Satu\nPembelian impulsif dibahas di sini.',
    'Lanjutan bab. Keterlibatan pelanggan dibahas di sini.',
    'Penutup bab menyebut Nuraiman dan Smith.',
    'Glosarium\nPembelian impulsif: definisi.',
    'INDEKS\nPembelian impulsif',
    'Daftar Pustaka\nNuraiman, S. (2024).',
]
GLOS = {'Pembelian impulsif': 'definisi',
        'Keterlibatan pelanggan': 'definisi',
        'Interaksi parasosial': 'definisi'}  # tanpa halaman isi -> dilewati
NAMA = ['Nuraiman', 'Smith']


def test_tanpa_indeks_nama():
    """Bentuk PRODUKSI: 'INDEKS' saja, tanpa 'Indeks Nama'."""
    d = tempfile.mkdtemp()
    p = _docx(os.path.join(d, 'buku.docx'),
              ['Sampul', 'DAFTAR ISI', 'BAB 1', 'Bab Satu', 'Glosarium',
               'INDEKS', 'Pembelian impulsif', 'Daftar Pustaka'])
    hasil = RI.isi_indeks(p, HALAMAN, GLOS, NAMA)
    teks = _teks(p)
    assert hasil['subjek'] == 2, hasil
    assert hasil['nama'] == 2, hasil
    assert 'Indeks Nama' in teks, 'seksi Indeks Nama harus DIBUAT'
    assert 'Pembelian impulsif .... 3' in teks, 'nomor dari isi, bukan TOC'
    assert 'Nuraiman .... 5' in teks
    shutil.rmtree(d)


def test_tanpa_seksi_indeks_sama_sekali():
    """Glosarium kosong: rakit tidak menulis seksi indeks sama sekali."""
    d = tempfile.mkdtemp()
    p = _docx(os.path.join(d, 'buku.docx'),
              ['Sampul', 'DAFTAR ISI', 'BAB 1', 'Bab Satu', 'Glosarium',
               'Daftar Pustaka'])
    hasil = RI.isi_indeks(p, HALAMAN, GLOS, NAMA)
    teks = _teks(p)
    assert 'Indeks Subjek' in teks and 'Indeks Nama' in teks
    assert 'Pembelian impulsif .... 3' in teks
    shutil.rmtree(d)


def test_seksi_sudah_ada():
    """Bentuk uji lama: kedua seksi ada + entri pageless -> diganti bernomor."""
    d = tempfile.mkdtemp()
    p = _docx(os.path.join(d, 'buku.docx'),
              ['Sampul', 'DAFTAR ISI', 'Pembelian impulsif .... ', 'BAB 1',
               'Bab Satu', 'Glosarium', 'Indeks Subjek', 'Pembelian impulsif',
               'Indeks Nama', 'Daftar Pustaka'])
    hasil = RI.isi_indeks(p, HALAMAN, GLOS, NAMA)
    teks = _teks(p)
    assert hasil['isi_awal'] == 3 and hasil['back_awal'] == 6, hasil
    assert 'Keterlibatan pelanggan .... 4' in teks
    assert 'Interaksi parasosial' not in teks, 'tanpa halaman -> dilewati'
    assert 'Smith .... 5' in teks
    shutil.rmtree(d)


def test_kumpulkan_istilah_menghapus_blok():
    d = tempfile.mkdtemp()
    os.makedirs(os.path.join(d, 'bab1', 'visual'))
    with open(os.path.join(d, 'bab1', 'bab1.txt'), 'w') as f:
        f.write('### 1.1 Definisi\nPembelian impulsif adalah pembelian.\n\n'
                '[ISTILAH BARU]\nPembelian impulsif: pembelian tanpa rencana.\n'
                '\n### 1.2 Perkembangan\nTeori berkembang.\n')
    glos, diambil = RI.kumpulkan_istilah(
        d, {1: {'title': 'Bab 1'}})
    assert diambil == 1 and glos['Pembelian impulsif'].startswith('pembelian')
    with open(os.path.join(d, 'bab1', 'bab1.txt')) as f:
        txt = f.read()
    assert '[ISTILAH BARU]' not in txt and '### 1.2 Perkembangan' in txt
    shutil.rmtree(d)


def test_kumpulkan_pustaka_dan_nama():
    d = tempfile.mkdtemp()
    for n, refs in ((1, ['Nuraiman, S. (2024). Impulse buying. Jurnal X.',
                         'Smith, J., & Lee, A. (2023). Parasocial. Jurnal Y.',
                         'Nuraiman, S. (2024). Impulse buying. Jurnal X.']),
                    (2, ['Wulandari, D. (2022). Engagement. Jurnal Z.'])):
        os.makedirs(os.path.join(d, f'bab{n}', 'visual'))
        with open(os.path.join(d, f'bab{n}', 'visual',
                               f'refs_bab{n}.json'), 'w') as f:
            json.dump(refs, f)
    pus = RI.kumpulkan_pustaka(d, {1: {'title': 'B1'}, 2: {'title': 'B2'}})
    assert len(pus) == 3, pus
    assert RI.kumpulkan_nama(pus) == ['Nuraiman', 'Smith', 'Lee', 'Wulandari']
    shutil.rmtree(d)


def main():
    tes = [v for k, v in sorted(globals().items())
           if k.startswith('test_') and callable(v)]
    gagal = 0
    for t in tes:
        try:
            t()
            print(f'PASS {t.__name__}')
        except Exception as e:
            gagal += 1
            print(f'FAIL {t.__name__}: {type(e).__name__}: {e}')
    print(f'{len(tes) - gagal}/{len(tes)} lulus')
    return 1 if gagal else 0


if __name__ == '__main__':
    raise SystemExit(main())

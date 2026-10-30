"""Uji G7 (koordinator): helper metadata, cap tier, filter singkatan.

Jalankan: /opt/data/.venv/bin/python3 uji_anti_gagal_g7.py
Tanpa DB produksi, tanpa order baru, tanpa LLM. Fixture sintetis saja.
"""
import os
import sys
import tempfile
import zipfile

sys.path.insert(0, '.')

LULUS = 0
GAGAL = []


def cek(nama, kondisi, detail=''):
    global LULUS
    if kondisi:
        LULUS += 1
        print('  LULUS %s' % nama)
    else:
        GAGAL.append(nama)
        print('  GAGAL %s %s' % (nama, detail))


def normalisasi_singkatan(hasil):
    """Terima dict token->kepanjangan atau list, kembalikan dict."""
    if isinstance(hasil, dict):
        return hasil
    out = {}
    for butir in (hasil or []):
        if isinstance(butir, dict):
            kunci = butir.get('singkatan') or butir.get('token') or butir.get('kode')
            nilai = butir.get('kepanjangan') or butir.get('arti') or ''
            if kunci:
                out[kunci] = nilai
        elif isinstance(butir, (list, tuple)) and len(butir) >= 2:
            out[butir[0]] = butir[1]
    return out


print('Uji G7 anti gagal (tier, metadata, singkatan)')
try:
    from docx import Document
    import buku_umum_assembly as ASM
    from tier_buku import boleh_ilustrasi
    from app.referensi_indeks import kumpulkan_singkatan
except ImportError as e:
    print('LEWAT: dependensi tidak tersedia: %s' % e)
    sys.exit(0)

# 1. Hitung kata: paragraf + sel tabel, tanpa ganda.
doc = Document()
doc.add_paragraph('satu dua tiga')
doc.add_paragraph('empat lima')
tab = doc.add_table(rows=1, cols=2)
tab.rows[0].cells[0].text = 'enam'
tab.rows[0].cells[1].text = 'tujuh delapan'
cek('hitung-kata-paragraf-tabel', ASM._hitung_kata_docx(doc) == 8,
    repr(ASM._hitung_kata_docx(doc)))

# 2. Pages metadata: tulis 42, tolak 0.
with tempfile.TemporaryDirectory() as tmp:
    jalur = os.path.join(tmp, 'uji.docx')
    doc.save(jalur)
    ASM.isi_halaman_docx(jalur, 42)
    with zipfile.ZipFile(jalur) as z:
        app_xml = z.read('docProps/app.xml').decode('utf-8', 'ignore')
    cek('pages-42-tertulis', '<Pages>42</Pages>' in app_xml, app_xml[:400])
    try:
        ASM.isi_halaman_docx(jalur, 0)
        cek('pages-0-ditolak', False, 'tidak raise')
    except ValueError:
        cek('pages-0-ditolak', True)

# 3. Cap tier: Hemat hanya bab 1 dan 2, Normal semua bab.
cek('tier-hemat-bab1', boleh_ilustrasi('hemat', 1) is True)
cek('tier-hemat-bab2', boleh_ilustrasi('hemat', 2) is True)
cek('tier-hemat-bab3-ditolak', boleh_ilustrasi('hemat', 3) is False)
cek('tier-hemat-bab5-ditolak', boleh_ilustrasi('hemat', 5) is False)
cek('tier-normal-bab5', boleh_ilustrasi('normal', 5) is True)

# 4. Singkatan: false positive hilang, kepanjangan derivasi menang.
with tempfile.TemporaryDirectory() as tmp:
    os.makedirs(os.path.join(tmp, 'bab1'))
    teks = (
        'Menurut RPIR et al., 2020, pasar berubah. SUB-BAB berikut membahas '
        'Usaha Mikro, Kecil, dan Menengah (UMKM) di daerah. '
        'Usaha Mikro, Kecil, dan Menengah adalah tulang punggung ekonomi. '
        'Resource Based View menjelaskan keunggulan. '
        'Resource Based View dipakai lagi di sini, dan RBV dirujuk. '
        'Technology Organization Environment menjadi kerangka. '
        'Technology Organization Environment dipakai lagi, dan TOE dirujuk. '
        'CATATAN EDITORIAL: bagian ini perlu dicek ulang.\n')
    with open(os.path.join(tmp, 'bab1', 'bab1.txt'), 'w', encoding='utf-8') as f:
        f.write(teks)
    sing = normalisasi_singkatan(kumpulkan_singkatan(tmp, {1: {'title': 'Bab 1'}}))
    cek('singkatan-RPIR-hilang', 'RPIR' not in sing, repr(sing))
    cek('singkatan-SUB-hilang', 'SUB' not in sing, repr(sing))
    cek('singkatan-CATATAN-hilang', 'CATATAN' not in sing, repr(sing))
    cek('singkatan-UMKM-kepanjangan',
        'Usaha Mikro' in (sing.get('UMKM') or ''), repr(sing))
    cek('singkatan-RBV-derivasi',
        'Resource Based View' in (sing.get('RBV') or ''), repr(sing))
    cek('singkatan-TOE-derivasi',
        'Technology Organization Environment' in (sing.get('TOE') or ''), repr(sing))

print('LULUS %d, GAGAL %d' % (LULUS, len(GAGAL)))
if GAGAL:
    print('GAGAL: %s' % ', '.join(GAGAL))
    sys.exit(1)
sys.exit(0)
